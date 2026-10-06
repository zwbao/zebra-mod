#!/usr/bin/env python3
"""Build zebra/data/china_rare_drug_approvals.json: rare-disease drug approvals and designations in China,
parsed only from official documents (MOF/STA/NMPA VAT lists, CDE urgently-needed overseas drug lists, CDE/NMPA
annual drug review reports, NMPA 创新药品名录 approval news).

How to rerun
------------
    python3 -I tools/china_access/build_drug_approvals.py --dl DIR            # parse files already in DIR
    python3 -I tools/china_access/build_drug_approvals.py --dl DIR --download # first re-fetch what a script can fetch

DIR holds the downloaded official files under the names listed in SOURCES / NEWS_DIR below. `--download` re-fetches
the gov.cn, chinatax.gov.cn, mof.gov.cn and nmpa.gov.cn static files (PDF/DOCX attachments) with urllib (proxy
honoured, >=1.2 s between requests). Pages on www.cde.org.cn and the HTML pages of www.nmpa.gov.cn answer scripts
with a JavaScript challenge (CDE: HTTP 202 + challenge page; NMPA: HTTP 412), so those files were saved from a real
browser session (ego-browser: page navigation, then in-page fetch of the attachment/article URL) and must already be
in DIR; the script stops with a list of missing files otherwise. Every source's sha256/bytes/retrieved_at (file
mtime, UTC) is recorded in the output's provenance.

Needs: python3 >= 3.9, openpyxl, pdfplumber; `pdftotext` (poppler); `soffice` (LibreOffice) to read the .doc
attachment of the CDE second-batch list; xlrd (optional, for the gov.cn .xls cross-check copy).
Output: zebra/data/china_rare_drug_approvals.json (override with --out).
"""
import argparse, datetime, glob, hashlib, html, json, os, re, shutil, subprocess, sys, tempfile, time
import unicodedata, urllib.request, zipfile
import xml.etree.ElementTree as ET

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
LISTS_JSON = os.path.join(REPO, 'zebra', 'data', 'china_rare_diseases.json')
OUT_DEFAULT = os.path.join(REPO, 'zebra', 'data', 'china_rare_drug_approvals.json')
UA = 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36'
NMPA_IMG = 'https://www.nmpa.gov.cn/directory/web/nmpa/images/'
CDE_NEWS = 'https://www.cde.org.cn/main/news/viewInfoCommon/'
CDE_ATT = 'https://www.cde.org.cn/main/att/download/'
NEWS_DIR = 'news'          # NMPA 创新药品名录 list pages (cxypxx_list_NN.html) and articles (art_<id>.html)
NEWS_BASE = 'https://www.nmpa.gov.cn/zhuanti/cxylqx/cxypxx/'

# ---------------------------------------------------------------- source registry
# files: {role: (local name, URL it was downloaded from, how: 'script' | 'browser')}
SOURCES = {
 'vat_2019': dict(
    title='财政部 海关总署 税务总局 药监局关于罕见病药品增值税政策的通知', document_no='财税〔2019〕24号',
    issuer='财政部 海关总署 税务总局 药监局', published='2019-02-20',
    origin='财政部网站 (gov.cn page states 来源：财政部网站)',
    files={'page': ('vat2019_govcn.html', 'https://www.gov.cn/zhengce/zhengceku/2019-10/15/content_5439874.htm', 'script'),
           'data': ('vat2019_govcn.xlsx', 'https://www.gov.cn/zhengce/zhengceku/2019-10/15/5439874/files/d2410003800741658b6694a36aa5cc56.xlsx', 'script'),
           'copy2_page': ('vat2019_fgk.html', 'https://fgk.chinatax.gov.cn/zcfgk/c102416/c5202340/content.html', 'script'),
           'copy2': ('vat2019_fgk.xlsx', 'https://fgk.chinatax.gov.cn/zcfgk/c102416/c5202340/5202340/files/%E7%BD%95%E8%A7%81%E7%97%85%E8%8D%AF%E5%93%81%E6%B8%85%E5%8D%95%EF%BC%88%E7%AC%AC%E4%B8%80%E6%89%B9%EF%BC%89.xlsx', 'script')}),
 'vat_2020': dict(
    title='财政部 海关总署 税务总局 药监局关于发布第二批适用增值税政策的抗癌药品和罕见病药品清单的公告',
    document_no='财政部 海关总署 税务总局 药监局公告2020年第39号', issuer='财政部 海关总署 税务总局 药监局', published='2020-09-30',
    files={'page': ('vat2020_mof.html', 'https://szs.mof.gov.cn/zhengcefabu/202009/t20200930_3598458.htm', 'script'),
           'data': ('vat2020_mof_list2.xlsx', 'https://szs.mof.gov.cn/zhengcefabu/202009/P020200930404323129068.xlsx', 'script'),
           'copy2_page': ('vat2020_govcn.html', 'https://www.gov.cn/zhengce/zhengceku/2020-10/09/content_5549951.htm', 'script'),
           'copy2': ('vat2020_govcn_list2.xls', 'https://www.gov.cn/zhengce/zhengceku/2020-10/09/5549951/files/bcabee97b7854d10aeb2b64a7a7b9a8f.xls', 'script'),
           'copy3_page': ('vat2020_fgk.html', 'https://fgk.chinatax.gov.cn/zcfgk/c102416/c5202182/content.html', 'script'),
           'copy3': ('vat2020_fgk_list2.xls', 'https://fgk.chinatax.gov.cn/zcfgk/c102416/c5202182/5202182/files/%E6%8A%97%E7%99%8C%E8%8D%AF%E5%93%81%E5%92%8C%E7%BD%95%E8%A7%81%E7%97%85%E8%8D%AF%E5%93%81%E6%B8%85%E5%8D%95%EF%BC%88%E7%AC%AC%E4%BA%8C%E6%89%B9%EF%BC%89.xls', 'script')}),
 'vat_2022': dict(
    title='关于发布第三批适用增值税政策的抗癌药品和罕见病药品清单的公告',
    document_no='财政部 海关总署 税务总局 药监局公告2022年第35号', issuer='财政部 海关总署 税务总局 药监局', published='2022-11-14',
    origin='税务总局网站 (gov.cn page states 来源：税务总局网站)',
    files={'page': ('vat2022_govcn.html', 'https://www.gov.cn/zhengce/zhengceku/2022-11/22/content_5728197.htm', 'script'),
           'data': ('vat2022_govcn_list3.pdf', 'https://www.gov.cn/zhengce/zhengceku/2022-11/22/5728197/files/73ae9c35788f40c1a5b3766409dde390.pdf', 'script'),
           'names': ('vat2022_govcn_names.pdf', 'https://www.gov.cn/zhengce/zhengceku/2022-11/22/5728197/files/0aeac5448c9d43989855585ebea96ef4.pdf', 'script'),
           'copy2_page': ('vat2022_fgk.html', 'https://fgk.chinatax.gov.cn/zcfgk/c102416/c5201990/content.html', 'script'),
           'copy2': ('vat2022_fgk_list3.pdf', 'https://fgk.chinatax.gov.cn/zcfgk/c102416/c5201990/5201990/files/%E6%8A%97%E7%99%8C%E8%8D%AF%E5%93%81%E5%92%8C%E7%BD%95%E8%A7%81%E7%97%85%E8%8D%AF%E5%93%81%E6%B8%85%E5%8D%95%EF%BC%88%E7%AC%AC%E4%B8%89%E6%89%B9%EF%BC%89.pdf', 'script')}),
 'mof_2026_10': dict(
    title='关于增值税法施行后增值税优惠政策衔接事项的公告', document_no='财政部 税务总局公告2026年第10号',
    issuer='财政部 税务总局', published='2026-01-30',
    files={'page': ('mof_2026_10.html', 'https://www.mof.gov.cn/jrttts/202602/t20260203_3983175.htm', 'script')}),
 'cde_list1': dict(
    title='关于发布第一批临床急需境外新药名单的通知', document_no=None, issuer='国家药品监督管理局药品审评中心', published='2018-11-01',
    files={'page': ('cdelist1.html', CDE_NEWS + '21de8acd6c395746b041b2ad93eb5c43', 'browser'),
           'data': ('cdelist1_att1.docx', CDE_ATT + '4ab933834f37673b660760d91ece9996', 'browser')}),
 'cde_list2': dict(
    title='关于发布第二批临床急需境外新药名单的通知', document_no=None, issuer='国家药品监督管理局药品审评中心', published='2019-05-29',
    files={'page': ('cdelist2.html', CDE_NEWS + '82f3bf94dc2c38d1a24d851f0e44914b', 'browser'),
           'data': ('cdelist2_att1.doc', CDE_ATT + '574b11bee371f0e239a06d34a4128726', 'browser')}),
 'cde_list3': dict(
    title='关于发布第三批临床急需境外新药名单的通知', document_no=None, issuer='国家药品监督管理局药品审评中心', published='2020-11-19',
    files={'page': ('cdelist3.html', CDE_NEWS + '08818b168ccc85db9a42a0f6623b5688', 'browser'),
           'data': ('cdelist3_att1.docx', CDE_ATT + 'fb5b9e0ed7773d275880c92a644f07de', 'browser')}),
 'rep_2018': dict(
    title='2018年度药品审评报告', document_no=None, issuer='国家药品监督管理局药品审评中心 (published by 国家药品监督管理局)', published='2019-07-01',
    files={'page': ('rep2018_nmpa.html', 'https://www.nmpa.gov.cn/xxgk/fgwj/gzwj/gzwjyp/20190701175801236.html', 'browser'),
           'prio': ('rep2018_nmpa_att4.docx', NMPA_IMG + 'uL28jMgMjAxOMTqyfPGwM2ouf21xNPFz8jJ88bA0qnGt8P7taUuZG9jeA==.docx', 'script'),
           'urgent': ('rep2018_nmpa_att6.docx', NMPA_IMG + 'uL28jUgtdrSu8X6wdm0sryx0Oi+s83i0MLSqbXEyfPGwMnzxfrH6b2LmRvY3g=.docx', 'script')}),
 'rep_2019': dict(
    title='2019年度药品审评报告', document_no=None, issuer='国家药品监督管理局药品审评中心 (published by 国家药品监督管理局)', published='2020-07-30',
    files={'page': ('rep2019_nmpa.html', 'https://www.nmpa.gov.cn/xxgk/fgwj/gzwj/gzwjyp/20200731114330106.html', 'browser'),
           'prio': ('rep2019_nmpa_att4.docx', NMPA_IMG + '1610416113512058121.docx', 'script'),
           'urgent': ('rep2019_nmpa_att5.docx', NMPA_IMG + '1610416118113032556.docx', 'script')}),
 'rep_2020': dict(
    title='2020年度药品审评报告', document_no=None, issuer='国家药品监督管理局药品审评中心 (published by 国家药品监督管理局)', published='2021-06-21',
    files={'page': ('rep2020_nmpa.html', 'https://www.nmpa.gov.cn/xxgk/fgwj/gzwj/gzwjyp/20210621142436183.html', 'browser'),
           'urgent': ('rep2020_nmpa_att7.docx', NMPA_IMG + '1624257995048013353.docx', 'script'),
           'copy2_page': ('rep2020_cde.html', CDE_NEWS + '876bb5300cce2d3a5cf4f68c97c8a631', 'browser')}),
 'rep_2021': dict(
    title='2021年度药品审评报告', document_no=None, issuer='国家药品监督管理局药品审评中心 (published by 国家药品监督管理局)', published='2022-06-01',
    files={'page': ('rep2021_nmpa.html', 'https://www.nmpa.gov.cn/xxgk/fgwj/gzwj/gzwjyp/20220601110541120.html', 'browser'),
           'urgent': ('rep2021_nmpa_att3.docx', NMPA_IMG + '1654068428217047310.docx', 'script')}),
 'rep_2023': dict(
    title='2023年度药品审评报告', document_no=None, issuer='国家药品监督管理局药品审评中心', published='2024-02-04',
    origin=CDE_NEWS + '9506710a7471174ab169e98b0bbb9e23 (CDE page dated 20240204; its PDF attachment download timed out in the browser, so the identical-title PDF on nmpa.gov.cn was used)',
    files={'data': ('rep2023_nmpa.pdf', NMPA_IMG + '1707039824076019627.pdf', 'script'),
           'cde_page': ('rep2023_cde.html', CDE_NEWS + '9506710a7471174ab169e98b0bbb9e23', 'browser')}),
 'rep_2024': dict(
    title='2024年度药品审评报告', document_no=None, issuer='国家药品监督管理局药品审评中心', published='2025-03-18',
    files={'page': ('rep2024_cde.html', CDE_NEWS + '54538c67b7e764fc51666567fc620241', 'browser'),
           'data': ('rep2024_cde.pdf', CDE_ATT + '0eee904628787acd38cef8447f7433a3', 'browser')}),
 'rep_2025': dict(
    title='2025年度药品审评报告', document_no=None, issuer='国家药品监督管理局药品审评中心', published='2026-05-13',
    files={'page': ('rep2025_cde.html', CDE_NEWS + 'b93aba16ec47467317057dca4aac8437', 'browser'),
           'data': ('rep2025_cde.pdf', CDE_ATT + 'c88a9b8aa86f0e9a7212ace4726f813c', 'browser')}),
}

NOT_RETRIEVED = [
 {'what': '2022年度药品审评报告 (CDE/NMPA)', 'url': 'https://www.nmpa.gov.cn/xxgk/fgwj/gzwj/gzwjyp/20230906163722146.html ; https://www.cde.org.cn/main/news/viewInfoCommon/849b5a642142fc00738aff200077db11',
  'reason': 'retrieved, but both official copies publish the report only as page images (JPG, no text layer, no attachment); not OCR\'d, so no 2022 rows or counts are taken from it'},
 {'what': '2023年度药品审评报告 PDF attachment on cde.org.cn', 'url': CDE_ATT + 'ad344d9ae98c0daa3590cfc873e9579a',
  'reason': 'in-browser download timed out after 60 s; the same report PDF from nmpa.gov.cn was used instead'},
 {'what': 'CDE 2018-08-08 征求意见 list of 48 境外已上市临床急需新药 (English names only)', 'url': None,
  'reason': 'not fetched: superseded by the 2018 report 附件5, which lists the same 48 with review status; the 8 "已在近期获批上市" drugs are not named in the 2018-11-01 notice'},
 {'what': 'NMPA approval announcements outside the 创新药品名录 column (e.g. 药品批准证明文件送达信息, generic/imported rare-disease drugs without a news item)', 'url': 'https://www.nmpa.gov.cn/',
  'reason': 'not searched systematically; this file is not a register of all NMPA approvals'},
]

SEARCHES = [
 'exa web search: "财政部 海关总署 税务总局 药监局 关于罕见病药品增值税政策的通知 财税〔2019〕24号 第一批罕见病药品清单"',
 'exa web search: "关于发布第二批适用增值税政策的抗癌药品和罕见病药品清单的公告 财政部 海关总署 税务总局 药监局"',
 'exa web search: "关于发布第三批适用增值税政策的抗癌药品和罕见病药品清单的公告"; "第四批适用增值税政策的抗癌药品和罕见病药品清单 公告 2023 2024 2025" (no 第四批 found; 财政部 税务总局公告2026年第10号 lists exactly three list documents as current)',
 'exa web search: "CDE 关于发布第一批/第二批/第三批临床急需境外新药名单的通知"; cde.org.cn site search (title) "临床急需境外新药名单", "药品审评报告"',
 'web search (gov.cn, cde.org.cn): "年度药品审评报告 药审中心 罕见病用药 附件"; nmpa.gov.cn 工作文件 list (gzwjyp) for 2018-2025 reports',
 'nmpa.gov.cn 创新药品名录 column (zhuanti/cxylqx/cxypxx, all list pages): every article read; kept when the text mentions 罕见 or the indication names a disease on the national lists',
]


STATUS_LAB = {2018: '附件5 第一批临床急需境外新药的审评审批情况', 2019: '附表5 境外已上市临床急需新药审评审批情况',
              2020: '附件7 2020年境外已上市临床急需新药审评审批情况', 2021: '附件3 2021年临床急需境外新药审评审批情况'}

# ---------------------------------------------------------------- helpers
def die(msg):
    sys.exit('ERROR: ' + msg)


def sha256(p):
    with open(p, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


def utc_iso(ts):
    return datetime.datetime.fromtimestamp(ts, datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')


def norm(s):
    """NFKC (full-width -> ASCII, Ⅷ -> VIII), drop all whitespace."""
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', s or ''))


def clean(s):
    return re.sub(r'\s+', ' ', (s or '').replace('　', ' ')).strip()


def html_text(path):
    raw = open(path, 'rb').read()
    t = raw.decode('utf-8', errors='replace')
    t = re.sub(r'<(script|style)\b.*?</\1>', '', t, flags=re.S | re.I)
    t = re.sub(r'<br\s*/?>|</p>|</div>|</tr>|</h\d>|</li>', '\n', t, flags=re.I)
    t = html.unescape(re.sub(r'<[^>]+>', ' ', t))
    t = re.sub(r'[ \t　\xa0]+', ' ', t)
    return re.sub(r'\s*\n\s*', '\n', t).strip()


W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def docx_tables(path):
    """[[row cells]] per table; a cell's paragraphs joined with '\n'."""
    root = ET.fromstring(zipfile.ZipFile(path).read('word/document.xml'))
    out = []
    for tbl in root.iter(W + 'tbl'):
        rows = []
        for tr in tbl.findall(W + 'tr'):
            cells = []
            for tc in tr.findall(W + 'tc'):
                paras = [''.join(t.text or '' for t in p.iter(W + 't')) for p in tc.iter(W + 'p')]
                cells.append('\n'.join(x for x in paras if x.strip()))
            rows.append(cells)
        out.append(rows)
    return out


def doc_to_docx(path, tmpdir):
    if not shutil.which('soffice'):
        die('soffice (LibreOffice) is needed to read ' + path)
    subprocess.run(['soffice', '--headless', '--convert-to', 'docx', '--outdir', tmpdir, path],
                   check=True, capture_output=True, timeout=180)
    out = os.path.join(tmpdir, os.path.splitext(os.path.basename(path))[0] + '.docx')
    if not os.path.exists(out):
        die('soffice did not convert ' + path)
    return out


def pdftext(path):
    return subprocess.run(['pdftotext', '-layout', path, '-'], check=True, capture_output=True).stdout.decode('utf-8')


def fetch(url, dest):
    req = urllib.request.Request(url, headers={'User-Agent': UA})
    with urllib.request.urlopen(req, timeout=90) as r:
        data = r.read()
        status = r.status
    if status != 200 or not data:
        raise RuntimeError(f'{url}: HTTP {status}, {len(data)} bytes')
    if data[:200].lstrip().lower().startswith((b'<!doctype', b'<html')) and not dest.endswith(('.html', '.htm')):
        raise RuntimeError(f'{url}: got an HTML page instead of a file (challenge?)')
    with open(dest, 'wb') as f:
        f.write(data)
    return len(data)


# ---------------------------------------------------------------- national lists + disease matching
def load_lists():
    d = json.load(open(LISTS_JSON, encoding='utf-8'))
    return d['diseases']


# a list term must not be read inside these longer words (documented precision guards)
BLOCK_BEFORE = {'血友病': ('血管性', '获得性'), '天疱疮': ('类',), '骨肉瘤': ('软',), '黑色素瘤': ()}


def disease_terms(diseases):
    """(list, no, list_name, term). Terms: full name; '/'-separated alternatives; for a parenthetical that is not
    a 型-qualifier: the name without it, and (if trailing) its content."""
    out = []
    for d in diseases:
        name = d['name_zh']
        alts = [a for a in re.split(r'/', name) if a.strip()]
        terms = {norm(name)}
        for a in alts:
            terms.add(norm(a))
            for m in re.finditer(r'[（(]([^（）()]*)[）)]', a):
                inner = m.group(1)
                if '型' in inner:
                    continue
                terms.add(norm(a[:m.start()] + a[m.end():]))
                if m.end() == len(a.rstrip()):
                    terms.add(norm(inner))
        for t in terms:
            if len(t) >= 3:
                out.append((d['list'], d['no'], name, t))
    return out


def match_diseases(text, terms):
    """[(list, no, list_name, term)] for list terms occurring verbatim (after NFKC/whitespace normalisation,
    case-insensitive) in text; a match lying inside a longer match of another disease is dropped."""
    if not text:
        return []
    t = norm(text).lower().replace('综合症', '综合征')
    hits = []
    for lst, no, name, term in terms:
        tl = term.lower()
        start = 0
        while True:
            i = t.find(tl, start)
            if i < 0:
                break
            before = t[max(0, i - 3):i]
            if not any(before.endswith(b) for b in BLOCK_BEFORE.get(term, ())):
                hits.append((i, i + len(tl), lst, no, name, term))
            start = i + 1
    keep = []
    for h in hits:
        inside = any(o is not h and (o[2], o[3]) != (h[2], h[3]) and o[0] <= h[0] and h[1] <= o[1] and (o[1] - o[0]) > (h[1] - h[0])
                     for o in hits)
        if not inside:
            keep.append(h)
    seen, res = set(), []
    for h in sorted(keep, key=lambda x: (x[2], x[3], -(x[1] - x[0]))):
        if (h[2], h[3]) in seen:
            continue
        seen.add((h[2], h[3]))
        res.append((h[2], h[3], h[4], h[5]))
    return res


# ---------------------------------------------------------------- name keys / merging
FORM_PREFIX = ('注射用', '吸入用', '口服用')
FORM_SUFFIX = sorted(['口服溶液用散', '口服混悬液', '口服液', '口服混悬剂', '口服溶液', '干混悬剂', '混悬液', '肠溶胶囊', '软胶囊',
                      '缓释胶囊', '胶囊', '缓释片', '分散片', '肠溶片', '口崩片', '片', '缓释注射液', '注射用浓溶液', '浓溶液',
                      '注射液', '滴眼液', '凝胶', '乳膏', '颗粒', '散', '微球', '吸入溶液', '溶液', '喷雾剂', '气雾剂', '贴剂', '栓'],
                     key=len, reverse=True)
TRAIL_PAREN = re.compile(r'\((预充式|皮下注射|静脉注射|预灌封)\)$')


def strip_form(n):
    n = TRAIL_PAREN.sub('', n)
    for p in FORM_PREFIX:
        if n.startswith(p) and len(n) > len(p) + 1:
            n = n[len(p):]
            break
    for s in FORM_SUFFIX:
        if n.endswith(s) and len(n) > len(s) + 1:
            n = n[:-len(s)]
            break
    return n


class Merger:
    def __init__(self):
        self.alias = {}      # stripped name -> canonical key
        self.alias_note = {}

    def add_alias(self, name, canon, note):
        k = strip_form(norm(name))
        c = strip_form(norm(canon))
        if k != c and k not in self.alias:
            self.alias[k] = c
            self.alias_note[k] = note

    def key(self, name):
        k = strip_form(norm(name))
        return self.alias.get(k, k)


# ---------------------------------------------------------------- parsers
def ev(sid, what, **kw):
    e = {'source_id': sid, 'what': what}
    e.update({k: v for k, v in kw.items() if v is not None})
    return e


def parse_vat2019(D, S):
    import openpyxl
    rows, prep, api = [], [], []
    ws = openpyxl.load_workbook(os.path.join(D, S['vat_2019']['files']['data'][0]), data_only=True).worksheets[0]
    section = None
    for r in ws.iter_rows(values_only=True):
        c0 = clean(str(r[0])) if r[0] is not None else ''
        if c0.startswith('一、罕见病药品制剂'):
            section = 'prep'; continue
        if c0.startswith('二、罕见病药品原料药'):
            section = 'api'; continue
        if not re.fullmatch(r'\d+', c0) or section is None:
            continue
        if section == 'prep':
            prep.append({'no': int(c0), 'ingredient': clean(str(r[1])), 'product': clean(str(r[2])), 'form': clean(str(r[3])), 'hs': str(r[4])})
        else:
            api.append({'no': int(c0), 'ingredient': clean(str(r[1])), 'hs': str(r[2])})
    assert [x['no'] for x in prep] == list(range(1, len(prep) + 1)) and [x['no'] for x in api] == list(range(1, len(api) + 1))
    return prep, api


def parse_vat2020_xlsx(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    return _vat2020_rows([[c for c in r] for r in ws.iter_rows(values_only=True)])


def parse_vat2020_xls(path):
    try:
        import xlrd
    except ImportError:
        return None
    sh = xlrd.open_workbook(path).sheet_by_index(0)
    rows = []
    for i in range(sh.nrows):
        row = []
        for c in sh.row_values(i):
            row.append(int(c) if isinstance(c, float) and c == int(c) else c)
        rows.append(row)
    return _vat2020_rows(rows)


def _vat2020_rows(rows):
    out, section = [], None
    for r in rows:
        r = list(r) + [None] * 5
        c0 = clean(str(r[0])) if r[0] not in (None, '') else ''
        if c0.startswith('二、罕见病药品制剂'):
            section = 'rare'; continue
        if c0.startswith(('一、', '（一）', '（二）')) and '罕见' not in c0:
            section = None if c0.startswith('一、') else section
            continue
        if section == 'rare' and re.fullmatch(r'\d+', c0):
            out.append({'no': int(c0), 'ingredient': clean(str(r[1])), 'product': clean(str(r[2])), 'form': clean(str(r[3])), 'hs': clean(str(r[4]))})
    assert [x['no'] for x in out] == list(range(1, len(out) + 1)), out
    return out


def parse_vat2022(path):
    t = pdftext(path)
    i = t.index('二、罕见病药品制剂和原料药')
    rare = t[i:]
    j = rare.index('（二）原料药')
    prep_t, api_t = rare[:j], rare[j:]
    prep = [{'no': int(m.group(1)), 'ingredient': clean(m.group(2)), 'form': m.group(3), 'hs': m.group(4)}
            for m in re.finditer(r'^\s*(\d+)\s+(\S.*?\S)\s{2,}(\S+)\s+(\d{8})\s*$', prep_t, re.M)]
    api = [{'no': int(m.group(1)), 'ingredient': clean(m.group(2)), 'hs': m.group(3)}
           for m in re.finditer(r'^\s*(\d+)\s+(\S.*?\S)\s{2,}(\d{8})\s*$', api_t, re.M)]
    assert [x['no'] for x in prep] == list(range(1, len(prep) + 1)) and [x['no'] for x in api] == list(range(1, len(api) + 1))
    return prep, api


def parse_vat2022_names(path):
    """附件3 部分药品活性成分通用名称情况: (main ingredient, listed name, which list)."""
    t = re.sub(r'\s+', '', pdftext(path))
    out = []
    for m in re.finditer(r'主要成份为“([^”]+)”的药品应视为与(.*?)序号(\d+)“([^”]+)”一致', t):
        out.append({'as_registered': m.group(1), 'listed_as': m.group(4), 'ref': m.group(2)[-40:] + '序号' + m.group(3)})
    return out


def parse_cde_list(path, tmp):
    if path.endswith('.doc'):
        path = doc_to_docx(path, tmp)
    tbl = max(docx_tables(path), key=len)
    head = [norm(c) for c in tbl[0]]
    col = lambda *names: next(i for i, h in enumerate(head) if any(n in h for n in names))
    ci = dict(no=col('序号'), name=col('药品名称'), co=col('企业名称'), ind=col('适应症'), why=col('列为临床急需原因'), date=col('批准日期'))
    out = []
    for r in tbl[1:]:
        if not r or not re.fullmatch(r'\d+', clean(r[ci['no']])):
            continue
        out.append({'no': int(clean(r[ci['no']])), 'name_en': clean(r[ci['name']].replace('\n', ' ')).rstrip(','),
                    'company': clean(r[ci['co']]), 'indication': clean(r[ci['ind']].replace('\n', ' ')), 'reason': clean(r[ci['why']].replace('\n', ' ')),
                    'date': norm_date(r[ci['date']])})
    assert [x['no'] for x in out] == list(range(1, len(out) + 1))
    return out


def norm_date(s):
    m = re.search(r'(\d{4})\s*[/\-.年]\s*(\d{1,2})\s*[/\-.月]\s*(\d{1,2})', s or '')
    return f'{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}' if m else None


def stated_number(text, pattern):
    m = re.search(pattern, text)
    return int(m.group(1)) if m else None


def en_key(s):
    """Key for matching English drug names across CDE documents: lower-case letters/digits of the INN part."""
    s = unicodedata.normalize('NFKC', s or '')
    m = re.search(r'\(([^()]*[A-Za-z][^()]*)\)', s)          # 'Brand (inn)' -> inn
    core = m.group(1) if m and not re.match(r'^\s*(recombinant|rdna)', m.group(1), re.I) else s
    core = re.sub(r'\b(injection|tablets?|capsules?|oral|delayed-release|granules|for|hcl|hydrochloride|mesylate|fumarate|sodium)\b', ' ', core, flags=re.I)
    core = re.sub(r'\d+(\.\d+)?\s*%|\d+\s*mg', ' ', core)
    return re.sub(r'[^a-z]', '', core.lower())


def parse_status_table(path, year):
    """2018 附件5 / 2019 附表5 / 2020 附件7 / 2021 附件3: CDE status of the urgently-needed overseas drugs."""
    rows = []
    for tbl in docx_tables(path):
        section = None
        head = None
        for r in tbl:
            cells = [clean(c.replace('\n', ' ')) for c in r]
            if len(cells) == 1 or (len(set(c for c in cells if c)) == 1 and not re.fullmatch(r'\d+', cells[0])):
                section = cells[0]           # 2021: 已批准品种 / 在审评品种 / 待申报品种
                continue
            if cells and norm(cells[0]) == '序号':
                head = [norm(c) for c in cells]
                continue
            if head is None:
                continue
            get = lambda *names: next((cells[i] for i, h in enumerate(head) if any(n in h for n in names) and i < len(cells)), '')
            name = get('药品名称', '境外药品名称', '药品通用名称')
            if not name:
                continue
            row = {'year': year, 'name': name, 'company': get('企业名称'), 'indication': get('适应症'),
                   'status': get('状态') or section or '', 'zh': get('通用名'), 'rare_col': get('罕见病目录'), 'date': norm_date(get('批准日期'))}
            if year == 2021:      # '依洛硫酸酯酶α注射液 （Elosulfase Alfa）' in one cell for approved drugs
                m = re.match(r'^(.*?)\s*[（(](.+)[）)]\s*$', name)
                if section and '已批准' in section and m:
                    row['zh'], row['name'] = clean(m.group(1)), clean(m.group(2))
                if section and '已批准' in section and not m:
                    row['zh'] = name
            rows.append(row)
    return rows


def report_narrative(text, year):
    """Numbered items of the report chapter on key approved products (重点品种 / 重点治疗领域品种)."""
    lines = text.split('\n')
    start = None
    for i, l in enumerate(lines):
        if re.search(r'审评通过的重点品种|重点治疗领域品种', l) and not re.search(r'\.{4,}', l):
            start = i
    if start is None:
        # 2019/2020: chapter heading not on its own line; start at the first category heading before item 1.
        for i, l in enumerate(lines):
            if re.match(r'^(抗肿瘤药物|新冠病毒疫苗和新冠肺炎治疗药物)：$', l.strip()):
                start = i; break
    if start is None:
        return []
    items, cat = [], None
    for l in lines[start + 1:]:
        s = l.strip()
        if re.match(r'^(第[一二三四五六七八九十]+章|[一二三四五六七八九十]+、|（[一二三四五六七八九十]+）)', s) and items:
            break
        m = re.match(r'^([^\d：:]{2,20})[：:]$', s)
        if m:
            cat = m.group(1); continue
        m = re.match(r'^(\d{1,3}(?:-\d{1,3})?)\s*[.．、]\s*(.+)$', s)
        if m:
            items.append({'n': m.group(1), 'cat': cat, 'text': m.group(2)})
    return items


def _spans(header):
    """Logical columns of a table header: (name, first raw index, last raw index). Some report tables come out
    of pdfplumber with every logical column split into several raw cells (only one carrying the header text);
    raw cells are assigned to the nearest header cell."""
    idx = [k for k, h in enumerate(header) if norm(h)]
    out = []
    for j, k in enumerate(idx):
        lo = 0 if j == 0 else (idx[j - 1] + k) // 2 + 1
        hi = len(header) - 1 if j == len(idx) - 1 else (k + idx[j + 1]) // 2
        out.append((norm(header[k]), lo, hi))
    return out


def pdf_appendices(path, wanted):
    """{appendix_no: (title, rows)} for numbered report appendices whose number is in `wanted`. Each row is a
    dict {logical column name: text}; continuation rows (empty 序号, e.g. across a page break) are appended to
    the previous row."""
    import pdfplumber
    res = {}
    with pdfplumber.open(path) as pdf:
        cur, spans = None, None
        for i, p in enumerate(pdf.pages):
            t = p.extract_text() or ''
            head = '\n'.join(t.split('\n')[:3])
            m = re.search(r'附[件表]\s*(\d+)\s*(20\d\d\s*年[^\n]*)', head)
            if m and i > 20:
                cur = int(m.group(1))
                if cur in wanted and cur not in res:
                    res[cur] = (clean(m.group(2)), [])
            if cur not in wanted:
                continue
            rows = res[cur][1]
            for tb in p.extract_tables():
                for r in tb:
                    raw = [clean((c or '').replace('\n', '')) for c in r]
                    if not any(raw):
                        continue
                    if '序号' in [norm(c) for c in raw]:
                        spans = _spans(raw)
                        continue
                    if spans is None:
                        continue
                    lg = {}
                    for name, lo, hi in spans:
                        lg[name] = ''.join(c for c in raw[lo:hi + 1] if c)
                    if re.fullmatch(r'\d+', lg.get('序号', '')):
                        rows.append(lg)
                    elif rows:
                        for k, v in lg.items():
                            if v and k != '序号':
                                rows[-1][k] = rows[-1].get(k, '') + v
    return res


def cell(row, *names):
    for k, v in row.items():
        if any(n in k for n in names):
            return v
    return ''


def parse_news(D):
    nd = os.path.join(D, NEWS_DIR)
    lists = sorted(glob.glob(os.path.join(nd, 'cxypxx_list_*.html')))
    if not lists:
        return None
    index, stated = {}, None
    for p in lists:
        raw = open(p, encoding='utf-8', errors='replace').read()
        m = re.search(r'共\s*([\d,]+)\s*条', raw)
        if m:
            stated = int(m.group(1).replace(',', ''))
        for m in re.finditer(r'<a\s[^>]*?href="[^"]*?cxypxx/(\d{14,})\.html"(.*?)</a>\s*(?:<span>\s*\(?(\d{4}-\d{2}-\d{2})\)?)?', raw, re.S):
            tm = re.search(r"title=(['\"])(.*?)\1", m.group(2), re.S)
            title = tm.group(2) if tm else m.group(2).split('>', 1)[-1]
            index.setdefault(m.group(1), (clean(html.unescape(re.sub(r'<[^>]+>', '', title))), m.group(3)))
    arts = []
    for aid in sorted(index):
        p = os.path.join(nd, f'art_{aid}.html')
        if not os.path.exists(p):
            arts.append({'id': aid, 'missing': True, 'title': index[aid][0]})
            continue
        t = html_text(p)
        title = index[aid][0]
        m = re.search(r'发布时间[:：]\s*(\d{4}-\d{2}-\d{2})', t)
        date = m.group(1) if m else index[aid][1]
        i = t.find('发布时间')
        body = t[i:] if i >= 0 else t
        body = re.sub(r'^发布时间[:：]\s*\d{4}-\d{2}-\d{2}\s*', '', body)
        body = body.split('本站由国家药品监督管理局主办')[0]
        body = re.sub(r'\n(网站声明|丨|网站使用指南|网站管理|-+地方药监局-+|-+直属单位-+).*', '', body, flags=re.S).strip()
        arts.append({'id': aid, 'title': title, 'date': date, 'body': body, 'path': p})
    return {'stated': stated, 'index': index, 'articles': arts, 'lists': lists}


def pathway(text):
    """Accelerated-pathway words as they appear in an official text, normalised to the programme names."""
    toks = []
    for w, lab in (('突破性治疗', '突破性治疗药物程序'), ('附条件批准', '附条件批准'), ('优先审评', '优先审评审批'), ('特别审批', '特别审批')):
        if w in (text or ''):
            toks.append(lab)
    return '; '.join(toks) or None


def news_fields(a):
    """(drug names, brand names, English names, indication sentence, pathway, exact approval date) of one
    创新药品名录 article, all as worded in the article."""
    title, body = a['title'], a['body']
    names = []
    m = re.search(r'(?:国家药监局|国家药品监督管理局)(?:有条件|附条件)?批准(?:中药创新药|1类创新药|创新药)?(.+?)(?:等\d+个品种)?上市', title) \
        or re.search(r'^用于.*?的(.+?)(?:获批)?上市$', title)
    if m:
        names = [clean(x) for x in re.split(r'[、和及]', m.group(1)) if clean(x)]
    if not names:   # older titles: '治疗罕见病药物注射用阿加糖酶β获批上市'
        m = re.search(r'^(.*?)(?:有条件)?获批上市$', title)
        if m:
            cand = re.split(r'药物|用药|新药|的', m.group(1))[-1].strip()
            if len(cand) >= 3:
                names = [cand]
    if not names:   # the product named in the text
        m = re.search(r'批准(?:[^。；]*?(?:申报的|研制的|生产的))?(?:1类创新药|1类新药|中药创新药|创新药)?'
                      r'([^，。、（(；“”]{2,40}?)(?:[（(](?:商品名|英文名)[:：][^）)]*[）)])?(?:上市|进口注册|的上市|注册申请)', body)
        if m:
            names = [clean(m.group(1).split('的')[-1])]
    names = [re.sub(r'生物类似药$', '', n) for n in names if n not in ('本品', '该药', '该药品')]
    brands = re.findall(r'商品名[:：]\s*([^）)，,]+)', body)
    en = re.findall(r'英文名[:：]\s*([^）)]+)', body)
    ind = None
    pat = r'((?:该[^，。；]{0,8}?)?(?:适应症为|适用于|用于|适用)[^。]+)。'
    m = re.search(pat, body.split('\n')[0]) or re.search(pat, body)
    if m:
        ind = clean(m.group(1))
    path = pathway(title + ' ' + body.split('。')[0])
    m = re.search(r'(20\d\d)年(\d{1,2})月(\d{1,2})日[，,]?\s*国家药(?:品监督管理局|监局)[^。]*?批准', body)
    exact = f'{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}' if m else None
    return names, brands, en, ind, path, exact


# ---------------------------------------------------------------- main build
def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--dl', required=True, help='directory with the downloaded official files')
    ap.add_argument('--download', action='store_true', help='re-fetch the script-fetchable files first')
    ap.add_argument('--out', default=OUT_DEFAULT)
    a = ap.parse_args()
    D = os.path.abspath(a.dl)
    t0 = time.time()

    if a.download:
        for sid, s in SOURCES.items():
            for role, (fn, url, how) in s['files'].items():
                if how != 'script':
                    continue
                n = fetch(url, os.path.join(D, fn))
                print(f'downloaded {fn} ({n} bytes)')
                time.sleep(1.2)
    missing = [f"{fn}  <-  {url}  ({how})" for s in SOURCES.values() for (fn, url, how) in s['files'].values()
               if not os.path.exists(os.path.join(D, fn))]
    if missing:
        die('missing files in --dl directory:\n  ' + '\n  '.join(missing))

    diseases = load_lists()
    terms = disease_terms(diseases)
    merger = Merger()
    recs = []           # one per (source mention)
    parts = {}          # source_id -> {what was parsed: number of rows}

    def part(sid, label, n):
        parts.setdefault(sid, {})[label] = n
    stated = {}
    cross = {}
    notes = {}
    tmp = tempfile.mkdtemp(prefix='zebra-nmpa-')
    P = lambda sid, role='data': os.path.join(D, SOURCES[sid]['files'][role][0])

    def add(sid, name_zh, what, kind, *, inn=None, inn_src=None, brand=None, ind=None, approved=None, year=None,
            ptype=None, date=None, announced=None, match_extra=(), rare_basis=None, en_name=None, status=None, scope=None):
        recs.append(dict(sid=sid, name_zh=name_zh, what=what, kind=kind, inn=inn, inn_src=inn_src, brand=brand, ind=ind,
                         approved=approved, year=year, ptype=ptype, date=date, announced=announced,
                         match_extra=list(match_extra), rare_basis=rare_basis, en_name=en_name, status=status, scope=scope))

    # ---- (a) VAT lists
    prep19, api19 = parse_vat2019(D, SOURCES)
    c19 = sha256(P('vat_2019')) == sha256(P('vat_2019', 'copy2'))
    cross['vat_2019'] = ('gov.cn policy-library xlsx vs chinatax 政策法规库 xlsx: ' + ('byte-identical (same sha256)' if c19 else 'DIFFERENT bytes'))
    for x in prep19:
        nm = x['ingredient'] if x['ingredient'] not in ('-', '') else x['product']
        if x['ingredient'] not in ('-', ''):
            merger.add_alias(x['product'], x['ingredient'], 'paired in 罕见病药品清单（第一批）(活性成分通用名称 / 药品名称)')
        add('vat_2019', nm, f"罕见病药品清单（第一批）一、制剂 序号{x['no']}: 活性成分 {x['ingredient']} / 药品名称 {x['product']} / 已获准上市的剂型 {x['form']}",
            'vat', approved=True)
    for x in api19:
        add('vat_2019', x['ingredient'], f"罕见病药品清单（第一批）二、原料药 序号{x['no']}: {x['ingredient']}", 'vat_api')
    part('vat_2019', '一、罕见病药品制剂', len(prep19)); part('vat_2019', '二、罕见病药品原料药', len(api19))
    notes['vat_2019'] = (f"{len(prep19)} 制剂 + {len(api19)} 原料药. §四: 本通知所称罕见病药品，是指经国家药品监督管理部门批准注册的罕见病药品制剂及原料药; "
                         "the list's column is 已获准上市的剂型, so each 制剂 row is taken as an official statement that the drug is approved in China (no date). "
                         "No indication is published. Where 活性成分通用名称 is '-', the 药品名称 is used as drug_zh.")

    rare20 = parse_vat2020_xlsx(P('vat_2020'))
    rare20_gov = parse_vat2020_xls(P('vat_2020', 'copy2'))
    same_xls = sha256(P('vat_2020', 'copy2')) == sha256(P('vat_2020', 'copy3'))
    if rare20_gov is None:
        cross['vat_2020'] = 'xlrd not installed: gov.cn .xls copy not parsed; ' + ('gov.cn and chinatax .xls copies byte-identical' if same_xls else 'gov.cn and chinatax .xls differ')
    else:
        diffs = [(a_['no'], a_['ingredient'], b_['ingredient']) for a_, b_ in zip(rare20, rare20_gov)
                 if (norm(a_['ingredient']), norm(a_['product'])) != (norm(b_['ingredient']), norm(b_['product']))]
        assert len(rare20) == len(rare20_gov), (len(rare20), len(rare20_gov))
        cross['vat_2020'] = (f"MOF xlsx (used) vs gov.cn .xls: {len(rare20)} vs {len(rare20_gov)} rare-disease rows, "
                             f"{'no differences in 活性成分/药品名称' if not diffs else 'differences: ' + str(diffs)}; gov.cn and chinatax .xls copies "
                             + ('byte-identical' if same_xls else 'differ'))
    for x in rare20:
        merger.add_alias(x['product'], x['ingredient'], 'paired in 抗癌药品和罕见病药品清单（第二批）(活性成分通用名称 / 药品名称)')
        add('vat_2020', x['ingredient'], f"抗癌药品和罕见病药品清单（第二批）二、罕见病药品制剂 序号{x['no']}: 活性成分 {x['ingredient']} / 药品名称 {x['product']} / 已获准上市的剂型 {x['form']}", 'vat', approved=True)
    part('vat_2020', '二、罕见病药品制剂', len(rare20))
    notes['vat_2020'] = f"{len(rare20)} 罕见病药品制剂 (the anti-cancer part of the same list is not used). No indication published."

    prep22, api22 = parse_vat2022(P('vat_2022'))
    c22 = sha256(P('vat_2022')) == sha256(P('vat_2022', 'copy2'))
    cross['vat_2022'] = 'gov.cn PDF vs chinatax 政策法规库 PDF: ' + ('byte-identical (same sha256)' if c22 else 'DIFFERENT bytes')
    eqv = parse_vat2022_names(P('vat_2022', 'names'))
    for e in eqv:
        merger.add_alias(e['as_registered'], e['listed_as'], '公告2022年第35号 附件3 部分药品活性成分通用名称情况: ' + e['ref'])
    for x in prep22:
        add('vat_2022', x['ingredient'], f"抗癌药品和罕见病药品清单（第三批）二、罕见病药品制剂 序号{x['no']}: {x['ingredient']} / 已获准上市的剂型 {x['form']}", 'vat', approved=True)
    for x in api22:
        add('vat_2022', x['ingredient'], f"抗癌药品和罕见病药品清单（第三批）二、（二）原料药 序号{x['no']}: {x['ingredient']}", 'vat_api')
    part('vat_2022', '二、（一）罕见病药品制剂', len(prep22)); part('vat_2022', '二、（二）罕见病药品原料药', len(api22))
    part('vat_2022', '附件3 name equivalences', len(eqv))
    notes['vat_2022'] = (f"{len(prep22)} 罕见病药品制剂 + {len(api22)} 原料药. §三: 各批清单中的抗癌药品和罕见病药品制剂需已获准上市; "
                         f"附件3 gives {len(eqv)} official name equivalences (used to merge, e.g. 盐酸芬戈莫德 = 芬戈莫德). No indication published.")

    mof = html_text(P('mof_2026_10', 'page'))
    vat_docs = re.findall(r'《([^《》]*罕见病药品[^《》]*)》', mof)
    assert any('财税〔2019〕24号' in mof for _ in [0]) and '2022年第35号' in mof and '2020年第39号' in mof
    later = [x for x in vat_docs if '第四批' in x or '第五批' in x]
    part('mof_2026_10', 'rare-disease list documents named as in force', len(vat_docs))
    notes['mof_2026_10'] = ('Not a list; used to check for later batches: its item 三(一)15 names the policies still in force — 财税〔2019〕24号 and the '
                            '第二批 (2020年第39号) and 第三批 (2022年第35号) list announcements — and no later rare-disease list'
                            + ('' if not later else f'; NOTE it also names: {later}') + '.')

    # ---- (b) CDE urgently-needed overseas drug lists
    cde = {}
    for k, sid in ((1, 'cde_list1'), (2, 'cde_list2'), (3, 'cde_list3')):
        rows = parse_cde_list(P(sid), tmp)
        page = html_text(P(sid, 'page'))
        n = stated_number(page, r'现将其他(\d+)个品种名单') or stated_number(page, r'等(\d+)个[^。]*?品种')
        stated[sid] = n
        assert n == len(rows), (sid, n, len(rows))
        cde[k] = rows
        part(sid, 'list rows', len(rows))
    notes['cde_list1'] = ('The notice says 48 drugs were selected, 8 of them "已在近期获批上市" (not named), and publishes the other 40. '
                          'The list gives English names (药品名称（活性成分）) and Chinese indications only, no Chinese drug name.')
    notes['cde_list2'] = 'Second batch: 26 drugs (English names, Chinese indications). Attachment is a .doc read via LibreOffice.'
    notes['cde_list3'] = 'Third batch: 7 drugs. The notice says the selection is now complete (no further batches).'

    # ---- (c) annual reports: status tables of the urgently-needed drugs
    st = {2018: parse_status_table(P('rep_2018', 'urgent'), 2018), 2019: parse_status_table(P('rep_2019', 'urgent'), 2019),
          2020: parse_status_table(P('rep_2020', 'urgent'), 2020), 2021: parse_status_table(P('rep_2021', 'urgent'), 2021)}
    rtext = {y: html_text(P(f'rep_{y}', 'page')) for y in (2018, 2019, 2020, 2021)}
    n18 = stated_number(rtext[2018], r'上述48个境外新药中已受理\d+个品种，(\d+)个品种已获批上市')
    a18 = sum('已批准' in r['status'] for r in st[2018])
    assert len(st[2018]) == 48 and n18 == a18, (len(st[2018]), n18, a18)
    n19 = stated_number(rtext[2019], r'目前已有(\d+)个品种批准上市或完成审评')
    m19 = re.search(r'目前已有(\d+)个品种批准上市或完成审评，(\d+)个品种正在进行技术审评，(\d+)个品种正在整理资料准备申报上市，'
                    r'(\d+)个品种正在整理资料且尚未提出注册申请，(\d+)个品种暂无申报上市计划，(\d+)个品种暂无法与持有企业取得联系', rtext[2019])
    t19 = sum(int(g) for g in m19.groups())
    assert len(st[2019]) == t19, (len(st[2019]), t19)
    a19 = sum(r['status'].strip() in ('已批准', '完成审评') for r in st[2019])
    n20 = stated_number(rtext[2020], r'已发布的三批(\d+)个品种临床急需')
    a20 = stated_number(rtext[2020], r'其中(\d+)个品种已获批上市或完成审评')
    assert len(st[2020]) == n20 and a20 == sum(r['status'] in ('已批准', '完成审评') for r in st[2020])
    a21 = stated_number(rtext[2021], r'(\d+)个品种获批上市，按审评时限')
    assert a21 == sum('已批准' in (r['status'] or '') for r in st[2021]) and len(st[2021]) == 81
    stated.update({'rep_2018:附件5': 48, 'rep_2020:附件7': n20, 'rep_2021:附件3': 81})
    for y, lab in ((2018, '附件5 urgent-drug status rows'), (2019, '附表5 urgent-drug status rows'), (2020, '附件7 urgent-drug status rows'),
                   (2021, '附件3 urgent-drug status rows')):
        part(f'rep_{y}', lab, len(st[y]))
    cross['rep_2018'] = f'附件5 has 48 rows = the 48 selected first-batch drugs; {a18} 已批准, equal to the report text ("{n18}个品种已获批上市")'
    stated['rep_2019:附表5'] = t19
    cross['rep_2019'] = (f'附表5 has {len(st[2019])} rows = the sum of the status counts in the text ({"+".join(m19.groups())}); text says {n19} 品种批准上市或完成审评, table has {a19} rows with 已批准/完成审评'
                         + ('' if n19 == a19 else ' (MISMATCH, recorded as published)'))
    cross['rep_2020'] = f'附件7 has {len(st[2020])} rows = "三批{n20}个品种"; {a20} 已获批上市或完成审评 in text = table'
    cross['rep_2021'] = f'附件3: {a21} rows under 已批准品种 = text "{a21}个品种获批上市"; 81 rows in total'

    # English-name join between the batch lists and the status tables
    class Look:
        """Find a drug in a report status table: by English-name key, else by a unique first-approval date."""
        def __init__(self, rows):
            self.rows, self.by = rows, {}
            for r in rows:
                self.by.setdefault(en_key(r['name']), r)

        def get(self, ek, date=None, who=None):
            r = self.by.get(ek)
            if r is None and date:
                c = [x for x in self.rows if x['date'] == date]
                r = c[0] if len(c) == 1 else None
                if r is not None and who:
                    date_joins.add((who, r['year'], r['name'], r['zh'] or '', date))
            return r

        def __contains__(self, ek):
            return ek in self.by
    date_joins = set()
    look = {y: Look(st[y]) for y in st}
    unmatched = []
    for k, rows in cde.items():
        for r in rows:
            for y in (2018, 2019, 2020, 2021):
                if y >= {1: 2018, 2: 2019, 3: 2020}[k]:
                    if look[y].get(en_key(r['name_en']), r['date']) is None:
                        unmatched.append((k, r['no'], r['name_en'], y))
    # 2018 附件5 drugs that are not in the published 40 (= the 8 "已在近期获批上市")
    first40 = {en_key(r['name_en']) for r in cde[1]}
    eight = [r for r in st[2018] if en_key(r['name']) not in first40]

    rare_first_col = {en_key(r['name']): r['rare_col'] for r in st[2018]}
    for k, sid in ((1, 'cde_list1'), (2, 'cde_list2'), (3, 'cde_list3')):
        for r in cde[k]:
            ek = en_key(r['name_en'])
            col = rare_first_col.get(ek) if k == 1 else None
            rare = ('罕见' in r['reason']) or (col and norm(col) != '非')
            # Chinese name / approval from the report status tables
            zh = None
            stat = []
            for y in (2018, 2019, 2020, 2021):
                s = look[y].get(ek, r['date'], who=r['name_en'])
                if not s:
                    continue
                if s['zh'] and s['zh'] not in ('—', '-'):
                    zh = s['zh']
                stat.append((y, s))
            add(sid, None, f"临床急需境外新药名单（第{'一二三'[k-1]}批）序号{r['no']}: {r['name_en']} — 列为临床急需原因: {r['reason'][:80]}",
                'cde', en_name=r['name_en'], inn=r['name_en'], inn_src=sid, ind=r['indication'],
                rare_basis=('reason/罕见病目录 column' if rare else None))
            recs[-1]['cde_key'] = ek
            recs[-1]['ind_src'] = sid
            for y, s in stat:
                sid_r = f'rep_{y}'
                lab = STATUS_LAB[y]
                status = s['status'].strip()
                approved = '已批准' in status
                what = f"{y}年度药品审评报告 {lab}: {s['name']} — 状态 {status} (as reported in the {y} report)" + (f"; 通用名 {s['zh']}" if s['zh'] and s['zh'] not in ('—', '-') else '')
                if y == 2018 and s['rare_col']:
                    what += f"; 第一批罕见病目录 column: {s['rare_col']}"
                inn2 = s['name'] if (y == 2021 and s['zh']) else None
                add(sid_r, None, what, 'cde_status', en_name=r['name_en'], approved=approved or None, status=status,
                    inn=inn2, inn_src=(sid_r if inn2 else None))
                recs[-1]['cde_key'] = ek
                recs[-1]['zh_from_status'] = s['zh'] if s['zh'] and s['zh'] not in ('—', '-') else None
    for s in eight:
        texts = bool(s['rare_col'] and norm(s['rare_col']) != '非')
        ek = en_key(s['name'])
        zh = None
        for y in (2019, 2020, 2021):
            s2 = look[y].get(ek, s['date'], who=s['name'])
            if s2 and s2['zh'] and s2['zh'] not in ('—', '-'):
                zh = s2['zh']
        add('rep_2018', None, f"2018年度药品审评报告 附件5: {s['name']} — 状态 {s['status'].strip()}; 第一批罕见病目录 column: {s['rare_col']} "
            "(one of the 48 selected first-batch drugs, not in the 40 published on 2018-11-01)", 'cde_status',
            en_name=s['name'], inn=s['name'], inn_src='rep_2018', approved=('已批准' in s['status']) or None, status=s['status'].strip(),
            rare_basis=('罕见病目录 column' if texts else None))
        recs[-1]['cde_key'] = ek
        recs[-1]['zh_from_status'] = zh
        for y in (2019, 2020, 2021):
            s2 = look[y].get(ek, s['date'])
            if s2:
                add(f'rep_{y}', None, f"{y}年度药品审评报告 {STATUS_LAB[y]}: {s2['name']} — 状态 {s2['status'].strip()} (as reported in the {y} report)" + (f"; 通用名 {s2['zh']}" if s2['zh'] and s2['zh'] not in ('—', '-') else ''),
                    'cde_status', en_name=s['name'], approved=('已批准' in s2['status']) or None, status=s2['status'].strip(),
                    inn=(s2['name'] if y == 2021 and s2['zh'] else None), inn_src=(f'rep_{y}' if y == 2021 and s2['zh'] else None))
                recs[-1]['cde_key'] = ek
                recs[-1]['zh_from_status'] = s2['zh'] if s2['zh'] and s2['zh'] not in ('—', '-') else None
                if y == 2020 or y == 2021:
                    recs[-1]['ind'] = s2['indication']; recs[-1]['ind_src'] = f'rep_{y}'

    # ---- (c) annual reports: priority-review lists with reason 罕见病 (2018 附件3, 2019 附表4)
    for y, role, lab in ((2018, 'prio', '附件3 2018年审评通过的优先审评药品名单'), (2019, 'prio', '附表4 2019年药审中心审评通过的优先审评品种')):
        n = tot = 0
        for tbl in docx_tables(P(f'rep_{y}', role)):
            for r in tbl:
                cells = [clean(c) for c in r]
                tot += bool(len(cells) >= 3 and re.fullmatch(r'\d+', cells[0]))
                if len(cells) >= 3 and re.fullmatch(r'\d+', cells[0]) and '罕见病' in cells[2]:
                    add(f'rep_{y}', cells[1], f"{y}年度药品审评报告 {lab} 序号{cells[0]}: {cells[1]} — 纳入优先审评的理由: {cells[2]} (审评通过)", 'prio')
                    n += 1
        part(f'rep_{y}', lab.split(' ')[0] + ' priority-review rows', tot)
        part(f'rep_{y}', lab.split(' ')[0] + ' rows with reason 罕见病 (used)', n)

    # ---- (c) annual reports: narrative chapter of key approved products (2018-2021 HTML)
    narr_counts = {}
    cross_2020 = None
    for y in (2018, 2019, 2020, 2021):
        items = report_narrative(rtext[y], y)
        kept = 0
        for it in items:
            if '-' in it['n']:
                continue
            m = re.match(r'^(.+?)[，,：:]', it['text'])
            if not m:
                continue
            name = clean(m.group(1))
            rest = it['text'][m.end():]
            first = clean(rest.split('。')[0])
            rare_kw = '罕见' in it['text'] or (it['cat'] or '').startswith('罕见病')
            approved = ('获批上市' in it['text']) or ('批准上市' in it['text'])
            add(f'rep_{y}', name, f"{y}年度药品审评报告 重点品种 {it['n']}. ({it['cat']}) {name}" + (" — text: 获批上市" if approved else ''),
                'narr', ind=first, approved=approved or None, rare_basis=('text mentions 罕见' if rare_kw else None))
            recs[-1]['ind_src'] = f'rep_{y}'
            kept += 1
        narr_counts[y] = kept
        part(f'rep_{y}', '重点品种 chapter items', kept)
        if y == 2020:
            cde_items = report_narrative(html_text(P('rep_2020', 'copy2_page')), 2020)
            a_ = [(i['n'], norm(i['text'])) for i in items]
            b_ = [(i['n'], norm(i['text'])) for i in cde_items]
            cross_2020 = ('2020 report: the 重点品种 chapter on nmpa.gov.cn (HTML, used) and on cde.org.cn (HTML) give '
                          + ('identical item lists' if a_ == b_ else f'different item lists ({len(a_)} vs {len(b_)} items)'))
    cross['rep_2020'] += '; ' + cross_2020

    # ---- (c) 2023-2025 PDF appendices
    pdf_counts = {}
    for y in (2023, 2024, 2025):
        sid = f'rep_{y}'
        txt = re.sub(r'\s+', '', pdftext(P(sid)))
        total = stated_number(txt, r'全年批准罕见病用药(\d+)个品种')
        prio = stated_number(txt, r'全年批准罕见病用药\d+个品种[^。]*?其中(\d+)个品种')
        stated[f'{sid}:罕见病用药(total)'] = total
        stated[f'{sid}:附件2'] = prio
        apps = pdf_appendices(P(sid), {1, 2, 4, 7})
        t2, rows2 = apps[2]
        assert '罕见病用药' in t2 and len(rows2) == prio, (y, t2, len(rows2), prio)
        n_inno = stated_number(txt, r'全年批准上市1类创新药(\d+)个品种')
        n_ext = stated_number(txt, r'(\d+)个品种，其中\d+个为(?:新批准上市|首次批准上市)[^。]*?附件4')
        m7 = re.search(r'共有(\d+)(?:个药品|件药品注册申请（(\d+)项适应症）)附条件批准上市', txt)
        n7 = int(m7.group(2) or m7.group(1)) if m7 else None
        for an, n_st in ((1, n_inno), (4, n_ext)):
            assert n_st == len(apps[an][1]) and [int(r['序号']) for r in apps[an][1]] == list(range(1, n_st + 1)), (y, an, n_st, len(apps[an][1]))
        stated.update({f'{sid}:附件1': n_inno, f'{sid}:附件4': n_ext, f'{sid}:附件7': n7})
        c7 = (f'附件7 has {len(apps[7][1])} numbered rows; text: {m7.group(0) if m7 else "?"}'
              + ('' if n7 == len(apps[7][1]) else ' (row count differs from the stated number; e.g. in 2024 the two combination '
                 'indications 贝莫苏拜单抗+安罗替尼 and 呋喹替尼+信迪利单抗 are listed under both partner drugs)' if y == 2024 else ' (row count differs from the stated number)'))
        cross[sid] = (f'附件2 {len(rows2)} rows = stated {prio} (全年批准罕见病用药{total}个品种，其中{prio}个通过优先审评审批); 附件1 {len(apps[1][1])} rows = stated 1类创新药 {n_inno}; '
                      f'附件4 {len(apps[4][1])} rows = stated {n_ext}; ' + c7)
        for an in (1, 2, 4, 7):
            part(sid, f'附件{an} rows', len(apps[an][1]))
        for r in rows2:
            name, ind = cell(r, '药品名称'), cell(r, '适应症')
            add(sid, name, f"{y}年度药品审评报告 附件2 {t2} 序号{cell(r, '序号')}: {name} ({cell(r, '上市许可持有人')})", 'rep_rare',
                ind=ind, approved=True, year=y, ptype='优先审评审批', rare_basis='appendix of rare-disease drugs')
            recs[-1]['ind_src'] = sid
        extra = 0
        for an in (1, 4, 7):
            if an not in apps:
                continue
            t_, rows_ = apps[an]
            if '批准' not in t_:
                continue
            for r in rows_:
                name = cell(r, '药品名称', '药品通用名称')
                ind = cell(r, '适应症', '附条件批准适应症', '功能主治')
                first = cell(r, '首次批准')
                mt = re.search(r'(首次批准上市|增加适应症|新增适应症)$', ind)
                if mt and an == 7:     # 附件7: the 首次批准上市/增加适应症 column is read into the indication cell
                    ind, first = ind[:mt.start()].rstrip('。 ') + '。', first or mt.group(1)
                if not name:
                    continue
                ptype = '附条件批准' if an == 7 else pathway(cell(r, '加快上市程序'))
                scope = 'new_indication' if re.search(r'新增|增加', first) and '首次' not in first else None
                add(sid, name, f"{y}年度药品审评报告 附件{an} {t_} 序号{cell(r, '序号')}: {name}" + (f" ({first})" if first else ''),
                    'rep_other', ind=ind, approved=True, year=y, ptype=ptype, scope=scope)
                recs[-1]['ind_src'] = sid
                extra += 1
        pdf_counts[y] = (len(rows2), extra)

    # ---- (d) NMPA 创新药品名录 news
    news = parse_news(D)
    news_sources = {}
    news_stats = None
    if news:
        arts = news['articles']
        got = [x for x in arts if not x.get('missing')]
        assert news['stated'] in (None, len(news['index'])), ('news index', news['stated'], len(news['index']))
        kept = 0
        for x in got:
            names, brands, en, ind, path, exact = news_fields(x)
            rare_kw = '罕见病' in x['body'] or '罕见病' in x['title']
            if not names:
                continue
            sid = f"nmpa_news_{x['id']}"
            news_sources[sid] = x
            brand = brands[0] if len(brands) == 1 and len(names) == 1 else None
            for nm in names:
                add(sid, nm, f"NMPA 创新药品名录 {x['date']}: {x['title']}", 'news', ind=ind, approved=True,
                    inn=(clean(en[0]) if len(en) == 1 and len(names) == 1 else None), inn_src=(sid if len(en) == 1 and len(names) == 1 else None),
                    year=int((exact or x['date'])[:4]), ptype=path, date=exact, announced=x['date'],
                    brand=brand, rare_basis=('text mentions 罕见病' if rare_kw else None))
                recs[-1]['ind_src'] = sid
            kept += 1
        news_stats = {'list_pages': len(news['lists']), 'stated_total': news['stated'], 'indexed': len(news['index']),
                      'downloaded': len(got), 'missing': [x['id'] for x in arts if x.get('missing')], 'articles_with_a_drug_name': kept}

    # ------------------------------------------------------------ merge
    # Chinese name for CDE-list rows: from the report status tables (latest wins)
    zh_by_cde = {}
    for r in recs:
        if r.get('cde_key') and r.get('zh_from_status'):
            zh_by_cde[r['cde_key']] = r['zh_from_status']
    groups, order = {}, []
    for r in recs:
        if r.get('cde_key'):
            zh = zh_by_cde.get(r['cde_key'])
            key = merger.key(zh) if zh else 'en:' + r['cde_key']
            r['name_zh'] = r['name_zh'] or zh
        else:
            key = merger.key(r['name_zh'])
        r['key'] = key
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(r)

    lists_by = {(d['list'], d['no']): d for d in diseases}
    src_rank = {'vat': 0, 'news': 1, 'rep_rare': 2, 'rep_other': 3, 'narr': 4, 'cde_status': 5, 'prio': 6, 'cde': 7, 'vat_api': 8}
    drugs = []
    for key in order:
        g = groups[key]
        rare_any = any(r['kind'] in ('vat', 'vat_api', 'rep_rare', 'prio') or r['rare_basis'] for r in g)
        # list links over every published indication text of the group
        links = {}
        for r in g:
            for text, s_ in ([(r['ind'], r.get('ind_src') or r['sid'])] if r['ind'] else []) + r['match_extra']:
                for lst, no, lname, term in match_diseases(text, terms):
                    if (lst, no) in links:
                        continue
                    how = f'indication ({s_}) contains "{term}"' + ('' if norm(lname) == term else f' (from list name "{lname}")')
                    links[(lst, no)] = how
        if not rare_any and not links:
            continue          # e.g. CDE urgent-list drugs for hepatitis C / psoriasis
        names = []
        for r in sorted(g, key=lambda r: src_rank[r['kind']]):
            if r['name_zh'] and r['name_zh'] not in names:
                names.append(r['name_zh'])
        vat_names = [r['name_zh'] for r in g if r['kind'] in ('vat', 'vat_api')]
        drug_zh = vat_names[0] if vat_names else (names[0] if names else None)
        inn_r = next((r for r in g if r['inn'] and r['inn_src'] and r['inn_src'].startswith('rep_2021')), None) or next((r for r in g if r['inn']), None)
        brand = next((r['brand'] for r in g if r['brand']), None)
        gi = [r for r in g if r['ind'] and r['scope'] != 'new_indication'] or [r for r in g if r['ind']]
        ind_r = (next((r for r in gi if r['kind'] == 'news'), None) or next((r for r in gi if r['kind'] in ('rep_rare', 'rep_other')), None)
                 or next((r for r in gi if r['kind'] == 'narr'), None) or next(iter(gi), None))
        claims = [r for r in g if r['approved']]
        approval = None
        if claims:
            # the claim that proves the earliest approval: its own date/year, else the publication date of its
            # document as an upper bound; within one year a claim that states the year wins.
            def bound(r):
                b = r['date'] or (f"{r['year']}-12-31" if r['year'] else None) or r['announced'] or SOURCES[r['sid']]['published']
                return (b[:4], r['year'] is None, b, src_rank[r['kind']])
            c = sorted(claims, key=bound)[0]
            if c['scope'] == 'new_indication':
                approval = {'date': None, 'year': None, 'type': None, 'source_id': c['sid'],
                            'note': f"the source reports an approval of a new indication in {c['year']}; the first approval date is not stated"}
            else:
                approval = {'date': c['date'], 'year': c['year'], 'type': c['ptype'], 'source_id': c['sid']}
            if c.get('announced'):
                approval['announced'] = c['announced']
        evid = []
        for r in g:
            e = ev(r['sid'], r['what'], status=r['status'])
            if r['approved']:
                e['approved'] = True
                if r['year']:
                    e['year'] = r['year']
                if r['scope']:
                    e['scope'] = r['scope']
                if r['ptype']:
                    e['type'] = r['ptype']
            if r['ind'] and ind_r is not None and r is not ind_r and norm(r['ind']) != norm(ind_r['ind']):
                e['indication_zh'] = r['ind']
            if e not in evid:
                evid.append(e)
        merged_via = sorted({merger.alias_note[k] for k in merger.alias if merger.alias[k] == key and any(strip_form(norm(r['name_zh'] or '')) == k for r in g)})
        row = {'drug_zh': drug_zh, 'names_zh': names if len(names) > 1 else None,
               'inn': inn_r['inn'] if inn_r else None, 'inn_source': inn_r['inn_src'] if inn_r else None,
               'brand': brand, 'indication_zh': ind_r['ind'] if ind_r else None, 'indication_en': None,
               'list_diseases': [{'list': k[0], 'no': k[1], 'name_zh': lists_by[k]['name_zh'], 'how': v} for k, v in sorted(links.items())],
               'approval': approval, 'evidence': evid}
        if merged_via:
            row['merged_via'] = merged_via
        if not drug_zh:
            last = max((r for r in g if r['kind'] == 'cde_status'), key=lambda r: r['sid'], default=None)
            row['note'] = ('No retrieved document gives a Chinese name for this CDE urgent-list drug; status '
                           + (f"'{last['status']}' in the {last['sid'][4:]} report" if last else 'unknown')
                           + '. A later approval under a Chinese name may appear as a separate row: no official document pairs the names.')
        drugs.append({k: v for k, v in row.items() if v is not None or k in ('drug_zh', 'inn', 'inn_source', 'brand', 'indication_zh', 'indication_en', 'approval')})

    drugs.sort(key=lambda d: (d['drug_zh'] is None, norm(d['drug_zh'] or d['inn'] or '')))

    # ------------------------------------------------------------ provenance
    used = {}
    for d_ in drugs:
        for e in d_['evidence']:
            used[e['source_id']] = used.get(e['source_id'], 0) + 1
    sources = {}
    for sid, s in SOURCES.items():
        fn, url, how = s['files'].get('data') or s['files']['page']
        p = os.path.join(D, fn)
        ent = {'title': s['title'], 'document_no': s.get('document_no'), 'issuer': s['issuer'], 'published': s['published'],
               'url': (s['files']['page'][1] if 'page' in s['files'] else url), 'origin': s.get('origin'),
               'attachment_url': url if 'data' in s['files'] else None, 'retrieved_at': utc_iso(os.path.getmtime(p)),
               'sha256': sha256(p), 'bytes': os.path.getsize(p),
               'rows': sum(v for k, v in parts.get(sid, {}).items() if '(used)' not in k and 'equivalences' not in k and 'named as in force' not in k),
               'parts': parts.get(sid), 'rows_used': used.get(sid, 0),
               'stated_total': None, 'cross_check': cross.get(sid), 'note': notes.get(sid)}
        if how == 'browser' or any(h == 'browser' for (_, _, h) in s['files'].values()):
            ent['retrieved_with'] = 'browser (site challenges scripts) for: ' + ', '.join(f for (f, _, h) in s['files'].values() if h == 'browser')
        extra_files = {role: {'file': f, 'url': u, 'sha256': sha256(os.path.join(D, f))} for role, (f, u, h) in s['files'].items() if role not in ('data',) and not (role == 'page' and 'data' not in s['files'])}
        if extra_files:
            ent['other_files'] = extra_files
        sources[sid] = {k: v for k, v in ent.items() if v is not None}
    sources['cde_list1']['stated_total'] = stated['cde_list1']
    sources['cde_list2']['stated_total'] = stated['cde_list2']
    sources['cde_list3']['stated_total'] = stated['cde_list3']
    for sid in ('cde_list1', 'cde_list2', 'cde_list3'):
        sources[sid]['cross_check'] = (f"every entry found by English name in the 2020 report 附件7 (81 = 48 + 26 + 7): "
                                        + ('yes' if not [u for u in unmatched if u[0] == int(sid[-1]) and u[3] == 2020] else 'NO, unmatched: ' + str([u for u in unmatched if u[0] == int(sid[-1]) and u[3] == 2020])))
    n19_rare = stated_number(rtext[2019], r'批准了(\d+)个用于治疗罕见病的、临床急需的药品')
    n20_rare = stated_number(rtext[2020], r'完成了(\d+)个用于治疗罕见病的')
    sources['rep_2018']['note'] = ('附件5 (48 first-batch urgent drugs with status and a 第一批罕见病目录 column), 附件3 (priority-review list; rows with reason 罕见病), '
                                   f'and the 重点品种 chapter of the HTML text. Stated: priority review included 63 applications for 儿童用药和罕见病用药 (no list).')
    sources['rep_2018']['stated_total'] = 48
    sources['rep_2019']['stated_total'] = stated['rep_2019:附表5']
    sources['rep_2019']['note'] = (f'附表5 ({len(st[2019])} urgent drugs, status and 通用名), 附表4 (priority-review list; reason 罕见病 rows), 重点品种 chapter. '
                                   f'Stated count: "{n19_rare}个用于治疗罕见病的、临床急需的药品" approved in 2019 (not listed by name).')
    sources['rep_2020']['note'] = (f'附件7 (81 urgent drugs, status and 通用名), 重点品种 chapter. Stated: '
                                   f'"完成了{n20_rare}个用于治疗罕见病的、临床急需的药品的技术审评" (not listed by name).')
    sources['rep_2020']['stated_total'] = n20
    sources['rep_2021']['note'] = f'附件3 (urgent drugs: 已批准 with Chinese name + active ingredient, 在审评, 待申报), 重点品种 chapter.'
    sources['rep_2021']['stated_total'] = 81
    for y in (2023, 2024, 2025):
        sid = f'rep_{y}'
        sources[sid]['stated_total'] = stated[f'{sid}:附件2']
        sources[sid]['note'] = (f'Stated: 全年批准罕见病用药 {stated[f"{sid}:罕见病用药(total)"]} 个品种, of which {stated[f"{sid}:附件2"]} via 优先审评审批 are listed by name in 附件2 '
                                f'(parsed {pdf_counts[y][0]}, asserted equal); the other {stated[f"{sid}:罕见病用药(total)"] - stated[f"{sid}:附件2"]} rare-disease approvals are not named. '
                                '附件1/4/7 (批准的创新药 / 境外已上市境内未上市 / 附条件批准) rows are used when the indication names a national-list disease or the drug is already in the file.')
    for sid, x in news_sources.items():
        if not used.get(sid):
            continue
        sources[sid] = {'title': x['title'], 'issuer': '国家药品监督管理局', 'published': x['date'], 'url': f"{NEWS_BASE}{x['id']}.html",
                        'retrieved_at': utc_iso(os.path.getmtime(x['path'])), 'sha256': sha256(x['path']), 'bytes': os.path.getsize(x['path']),
                        'rows': 1, 'rows_used': used.get(sid, 0), 'retrieved_with': 'browser (nmpa.gov.cn returns HTTP 412 to scripts)'}

    counts = {
        'drugs': len(drugs),
        'with_approval': sum(1 for d in drugs if d['approval']),
        'with_approval_year': sum(1 for d in drugs if d['approval'] and d['approval']['year']),
        'with_approval_date': sum(1 for d in drugs if d['approval'] and d['approval']['date']),
        'not_approved_per_sources': sum(1 for d in drugs if not d['approval']),
        'with_list_diseases': sum(1 for d in drugs if d['list_diseases']),
        'with_inn': sum(1 for d in drugs if d['inn']),
        'without_drug_zh': sum(1 for d in drugs if not d['drug_zh']),
        'distinct_list_diseases_linked': len({(x['list'], x['no']) for d in drugs for x in d['list_diseases']}),
        'source_mentions': len(recs),
        'vat_rare_preparations': {'2019': len(prep19), '2020': len(rare20), '2022': len(prep22)},
        'vat_rare_apis': {'2019': len(api19), '2022': len(api22)},
        'cde_urgent_lists': {'1': len(cde[1]), '2': len(cde[2]), '3': len(cde[3])},
        'report_rare_appendix_rows': {str(y): pdf_counts[y][0] for y in pdf_counts},
        'report_stated_rare_approvals': {str(y): stated[f'rep_{y}:罕见病用药(total)'] for y in (2023, 2024, 2025)},
        'nmpa_news': news_stats,
        'cde_status_joined_by_date': len(date_joins),
        'nmpa_news_articles_used': sum(1 for k in used if k.startswith('nmpa_news_')),
    }
    coverage = (
        'Covers: (1) every drug on the three official VAT 罕见病药品清单 (2019 第一批 21 制剂 + 4 原料药, 2020 第二批 14, 2022 第三批 19 + 1 原料药; '
        'a fourth batch was not found and the 2026 MOF/STA notice lists only these three) — these are NMPA-approved by definition but carry no '
        'indication, date or pathway; (2) every drug on the three CDE 临床急需境外新药名单 that the list or the 2018 report marks as a rare-disease drug '
        'or whose indication names a national-list disease, with the CDE review status reported in the 2018-2021 annual reports (approval only where '
        'a report says 已批准/获批); (3) the rare-disease drugs named in the 2018-2021 and 2023-2025 年度药品审评报告 (2023-2025: only the ones approved '
        'via 优先审评审批 are named — e.g. 15 of 45 in 2023, 20 of 55 in 2024, 12 of 48 in 2025 — plus 创新药/境外已上市/附条件批准 rows whose indication '
        'names a list disease); (4) NMPA 创新药品名录 approval news (2018-2026) that mention 罕见病 or whose indication names a list disease. '
        'NOT covered: this is not a register of NMPA approvals — most rare-disease approvals (generics, imported originals without a news item, the '
        'unnamed majority of each year\'s rare-disease approvals, the whole 2022 report which is published only as images) are missing, so a drug '
        'absent here may still be approved in China; approval rows rarely carry an exact date (news dates are announcement dates); indications '
        'are as published for that product/approval and may be narrower or broader than the current label; reimbursement (医保) status is not '
        'covered; list_diseases links are literal name matches only, so a drug can treat a listed disease without being linked. Rows without '
        'drug_zh are CDE urgent-list drugs known only by their English name, with CDE status up to the 2021 report; several were approved later '
        'under a Chinese name that appears here as a separate row (no official document pairs the names), so their approval: null means only '
        '"no approval found under this name".')
    out = {
        'schema': 1,
        'title': 'Rare-disease drug approvals and designations in China (NMPA/CDE/MOF official documents)',
        'provenance': {
            'built_at': datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'),
            'builder': 'tools/china_access/build_drug_approvals.py',
            'sources': sources,
            'counts': counts,
            'coverage': coverage,
            'not_retrieved': NOT_RETRIEVED + ([{'what': f'NMPA 创新药品名录 article {i}', 'url': f'{NEWS_BASE}{i}.html', 'reason': 'not downloaded'} for i in (news_stats or {}).get('missing', [])]),
            'method': {
                'merge': ('one row per drug: names are NFKC-normalised, a fixed list of dosage-form affixes is removed (注射用…, …注射液/片/胶囊/口服溶液…), '
                          'and names are joined when an official document pairs them (VAT list 活性成分↔药品名称; 2022 附件3 name equivalences); salts are '
                          'not stripped otherwise. CDE urgent-list drugs (English only) are joined to Chinese names through the report tables that give both.'),
                'approval': ('approval is non-null only when a retrieved official document says approved/获批/已获准上市; the row shows the claim with the '
                             'most precise date/year; every claim stays in evidence (approved: true). VAT 制剂 rows count as approved without a date.'),
                'list_diseases': ('a national-list name (or a "/"-alternative, or a non-型 parenthetical variant) occurring verbatim in a published indication '
                                  '(NFKC, whitespace and case ignored, 综合症 read as 综合征); a match inside a longer matched list name is dropped; 血友病 is not read '
                                  'inside 血管性血友病/获得性血友病, 天疱疮 not inside 类天疱疮, 骨肉瘤 not inside 软骨肉瘤. The 2018 report\'s 第一批罕见病目录 column '
                                  'is kept in evidence but not used for links (it is not an indication).'),
                'cde_status_join': ('CDE urgent-list drugs are found in the 2018-2021 report status tables by English name, else by a unique first '
                                    'foreign approval date; the date-based joins are listed here: ' + '; '.join(
                                        f'{a_} -> {y_} table "{n_}"' + (f' ({z_})' if z_ else '') + f' [{d_}]' for a_, y_, n_, z_, d_ in sorted(date_joins))),
                'inclusion': ('VAT rare-disease rows and report rare-disease appendices: all; CDE urgent lists: when the list\'s 列为临床急需原因 mentions '
                              '罕见 or the 2018 report marks a 第一批罕见病目录 disease, or the indication names a list disease. Report 重点品种 items, 2023-2025 '
                              '附件1/4/7 rows and NMPA 创新药品名录 news are parsed in full and kept when the text mentions 罕见 (news: 罕见病), when the indication '
                              'names a list disease, or when they name (after the merge rules) a drug that is already in the file from another source — so they '
                              'also add approval years/pathways to drugs found in the VAT and CDE lists.'),
                'searches': SEARCHES,
            },
        },
        'drugs': drugs,
    }
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    with open(a.out, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
        f.write('\n')
    shutil.rmtree(tmp, ignore_errors=True)
    print(f"wrote {a.out}: {len(drugs)} drugs, {counts['with_approval']} with approval, {counts['with_list_diseases']} linked to list diseases; "
          f"{time.time() - t0:.1f} s")
    if unmatched:
        print('CDE list names not found in report tables:', unmatched)


if __name__ == '__main__':
    main()
